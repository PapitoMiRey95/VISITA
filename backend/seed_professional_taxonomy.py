"""VIen Professional Construct — Phase 2 taxonomy seeding/import (PREVIEW).

Idempotent. Rerunning cannot create duplicates (every doc is upserted by its
stable `id`). Never runs at server startup / import — invoke explicitly:

    python seed_professional_taxonomy.py

Sources (authoritative Access exports, bundled under ./taxonomy_sources):
  Sphere (hard list 1=Health, 2=Non-Health; SphereID 3 = error, ignored),
  tblProfessionalArea, tblProfessionalSpeciality, tblProfessionalCredentials,
  tblLanguage.

Source IDs are PRESERVED exactly (never regenerated / normalized / merged /
renamed / repaired). Known legacy exceptions (Book Keeper 155, Paralegal 193,
Portuguese 103/104, duplicate credential labels) are imported as-is.

The three NEW VIen-controlled taxonomies (Areas of Practice, Practice Types,
Ontario Primary Care Models) use stable VIen IDs (vaop-/vpt-/vpcm-) and never
reuse Access numbering.
"""
import asyncio
import os
from pathlib import Path

import openpyxl
from motor.motor_asyncio import AsyncIOMotorClient

SRC = Path(__file__).parent / "taxonomy_sources"

FAMILY_MEDICINE_SPECIALTY_ID = "133"  # resolved from tblProfessionalSpeciality (AreaID 36)

SPHERES = [
    {"id": "1", "sphere_id": 1, "name": "Health"},
    {"id": "2", "sphere_id": 2, "name": "Non-Health"},
]

# --- NEW VIen taxonomy: Family Medicine Areas of Practice (97) ---
AOP_CATEGORIES = [
    ("GENERAL / COMPREHENSIVE", [
        "Comprehensive Family Medicine", "General Primary Care",
        "Preventive Care / Health Promotion", "Chronic Disease Management",
        "Complex / Multimorbidity Care"]),
    ("AGE / POPULATION", [
        "Newborn Care", "Child / Pediatric Health", "Adolescent / Teen Health",
        "Young Adult Health", "Adult Medicine", "Geriatric Care / Care of the Elderly",
        "Women\u2019s Health", "Men\u2019s Health", "LGBTQ2S+ Health",
        "Transgender / Gender-Affirming Primary Care", "Indigenous Health",
        "Newcomer / Immigrant Health", "Refugee Health", "Street Health / Homeless Health",
        "Vulnerable / Marginalized Populations", "Rural Medicine", "Remote / Northern Medicine",
        "Disability Health", "Developmental / Intellectual Disability Care"]),
    ("WOMEN / REPRODUCTIVE / MATERNITY", [
        "Reproductive Health", "Contraception", "Sexual Health", "STI Care", "Prenatal Care",
        "Low-Risk Obstetrics / Maternity Care", "Postpartum Care",
        "Breastfeeding / Lactation Care", "Menopause Care",
        "Cervical Screening / Women\u2019s Preventive Health"]),
    ("MENTAL HEALTH / SUBSTANCE USE", [
        "Mental Health", "Depression / Anxiety", "Serious Mental Illness",
        "Addiction Medicine", "Substance Use Disorders", "Opioid Agonist Therapy",
        "Alcohol Use Disorders", "Smoking / Tobacco Cessation", "Eating Disorders"]),
    ("CHRONIC / CLINICAL FOCUS", [
        "Diabetes Care", "Hypertension / Cardiovascular Risk",
        "Respiratory Care / Asthma / COPD", "Obesity / Weight Management",
        "Lipid Management", "Thyroid / Endocrine Primary Care",
        "Kidney Disease Primary Care", "Liver Disease Primary Care",
        "Arthritis / Rheumatologic Primary Care", "Neurologic Primary Care",
        "Dementia / Cognitive Care", "Frailty Care", "Osteoporosis",
        "Cancer Survivorship", "HIV Primary Care", "Hepatitis Care",
        "Infectious Disease Primary Care", "Dermatology in Primary Care", "Sleep Health"]),
    ("PAIN / MUSCULOSKELETAL", [
        "Chronic Pain", "Musculoskeletal Medicine", "Sports & Exercise Medicine",
        "Joint / Soft Tissue Injections", "Workplace / Occupational Injury Care"]),
    ("ACUTE / PROCEDURAL", [
        "Acute Care", "Urgent Care", "Emergency Medicine", "Minor Procedures",
        "Minor Surgery", "Skin Procedures / Dermatologic Procedures", "Suturing / Wound Care",
        "Cryotherapy", "Biopsy / Excision", "IUD Insertion / Removal",
        "Contraceptive Implant Procedures", "Joint Injection", "Trigger Point Injection",
        "Wound Care"]),
    ("OLDER ADULT / END-OF-LIFE", [
        "Palliative Care", "End-of-Life Care", "Long-Term Care", "Nursing Home Care",
        "Home-Based Care / House Calls"]),
    ("OTHER FOCUSED FAMILY MEDICINE", [
        "Travel Medicine", "Occupational Medicine", "Student Health", "School Health",
        "Correctional Medicine", "Shelter Medicine", "Sexual Medicine",
        "Hospital Medicine / Hospitalist Care", "Family Practice Anesthesia",
        "Virtual Care / Telemedicine", "Focused Family Medicine Practice"]),
]

PRACTICE_TYPES = [
    "Solo Practice", "Group Family Practice", "Walk-In Clinic", "Family Practice + Walk-In",
    "Community-Based Family Practice", "Family Health Team", "Community Health Centre",
    "Academic / Teaching Practice", "Hospital-Based Practice", "Hospitalist Practice",
    "Urgent Care Practice", "Rural Practice", "Remote / Northern Practice",
    "Long-Term Care Practice", "Nursing Home Practice", "Home-Based / House Call Practice",
    "Shelter / Street Health Practice", "Indigenous Primary Care", "Student / Campus Health",
    "Occupational Health Practice", "Virtual / Telemedicine Practice", "Locum Practice",
    "Focused Family Medicine Practice", "Interprofessional / Multidisciplinary Practice",
]

PRIMARY_CARE_MODELS = [
    ("Traditional Fee-for-Service", None),
    ("Comprehensive Care Model (CCM)", "CCM"),
    ("Family Health Group (FHG)", "FHG"),
    ("Family Health Network (FHN)", "FHN"),
    ("Family Health Organization (FHO)", "FHO"),
    ("Family Health Team (FHT)", "FHT"),
    ("Community Health Centre (CHC)", "CHC"),
    ("Blended Salary Model (BSM)", "BSM"),
    ("Nurse Practitioner-Led Clinic (NPLC)", "NPLC"),
    ("Rural and Northern Physician Group Agreement (RNPGA)", "RNPGA"),
    ("Group Health Centre (GHC)", "GHC"),
    ("Indigenous Primary Health Care Organization (IPHCO)", "IPHCO"),
    ("Other / Custom", None),
]


def _rows(fname):
    wb = openpyxl.load_workbook(SRC / fname, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    data = [r for r in rows[1:] if any(c is not None and str(c).strip() != "" for c in r)]
    return header, data


async def _upsert(coll, docs):
    for d in docs:
        await coll.update_one({"id": d["id"]}, {"$set": d}, upsert=True)


def build_areas():
    _, data = _rows("tblProfessionalArea.xlsx")
    out = []
    for AreaID, SphereID, Area, Required, Favourite, Favourite2 in data:
        if SphereID not in (1, 2):  # SphereID 3 = error placeholders (ZZZ/Z) -> ignore
            continue
        out.append({
            "id": str(AreaID), "area_id": AreaID, "sphere_id": str(SphereID),
            "name": Area, "required": bool(Required),
            "favourite": bool(Favourite), "favourite2": bool(Favourite2),
        })
    return out


def build_specialties():
    _, data = _rows("tblProfessionalSpeciality.xlsx")
    out = []
    for SpecialityID, AreaID, Speciality, SpecialityReturned, Favourite in data:
        out.append({
            "id": str(SpecialityID), "specialty_id": SpecialityID,
            "area_id": str(AreaID) if AreaID is not None else None,
            "speciality": Speciality, "speciality_returned": SpecialityReturned,
            "favourite": bool(Favourite),
        })
    return out


def build_credentials():
    _, data = _rows("tblProfessionalCredentials.xlsx")
    out = []
    for CredentialsID, Credentials, CredentialsReturned, Desc, Favourite in data:
        out.append({
            "id": str(CredentialsID), "credential_id": CredentialsID,
            "credentials": Credentials, "credentials_returned": CredentialsReturned,
            "desc": Desc, "favourite": bool(Favourite),
        })
    return out


def build_languages():
    _, data = _rows("tblLanguage.xlsx")
    out = []
    for LanguageID, Language, Favourite in data:
        out.append({
            "id": str(LanguageID), "language_id": LanguageID,
            "language": Language, "favourite": bool(Favourite),
        })
    return out


def build_aop():
    out = []
    n = 0
    for category, names in AOP_CATEGORIES:
        for name in names:
            n += 1
            out.append({
                "id": f"vaop-{n:03d}", "specialty_id": FAMILY_MEDICINE_SPECIALTY_ID,
                "category": category, "name": name, "active": True, "sort_order": n,
            })
    return out


def build_practice_types():
    return [{"id": f"vpt-{i:02d}", "name": name, "active": True, "sort_order": i}
            for i, name in enumerate(PRACTICE_TYPES, 1)]


def build_models():
    return [{"id": f"vpcm-{i:02d}", "name": name, "abbreviation": abbr,
             "active": True, "sort_order": i}
            for i, (name, abbr) in enumerate(PRIMARY_CARE_MODELS, 1)]


async def main():
    client = AsyncIOMotorClient(os.environ["MONGO_URL"])
    db = client[os.environ["DB_NAME"]]

    spheres = SPHERES
    areas = build_areas()
    specialties = build_specialties()
    credentials = build_credentials()
    languages = build_languages()
    aop = build_aop()
    practice_types = build_practice_types()
    models = build_models()

    await _upsert(db.professional_spheres, spheres)
    await _upsert(db.professional_areas, areas)
    await _upsert(db.professional_specialties, specialties)
    await _upsert(db.professional_credentials, credentials)
    await _upsert(db.professional_languages, languages)
    await _upsert(db.professional_areas_of_practice, aop)
    await _upsert(db.professional_practice_types, practice_types)
    await _upsert(db.professional_primary_care_models, models)

    print("=== VIen Professional taxonomy seed (idempotent upsert) — PREVIEW ===")
    print(f"Family Medicine SpecialityID: {FAMILY_MEDICINE_SPECIALTY_ID}")
    for coll, label in [
        (db.professional_spheres, "spheres"),
        (db.professional_areas, "areas (sphere 1/2 only)"),
        (db.professional_specialties, "specialties"),
        (db.professional_credentials, "credentials"),
        (db.professional_languages, "languages"),
        (db.professional_areas_of_practice, "areas_of_practice (Family Medicine)"),
        (db.professional_practice_types, "practice_types"),
        (db.professional_primary_care_models, "primary_care_models"),
    ]:
        print(f"  {label}: {await coll.count_documents({})}")
    print("Profiles untouched:", await db.professional_profiles.count_documents({}))
    client.close()


if __name__ == "__main__":
    asyncio.run(main())
